#!/usr/bin/env python3
"""Extrai da Tabela TACO (4ª edição, NEPA/UNICAMP) os valores por 100 g usados pelo app e grava data/taco.json.

Uso: npm run data:taco   (baixa a planilha oficial se não estiver em data/.cache/)

Fonte: Tabela Brasileira de Composição de Alimentos - TACO, 4ª edição revisada e ampliada.
NEPA/UNICAMP, Campinas, 2011. https://nepa.unicamp.br/
"""
import hashlib, json, sys, urllib.request
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
URL = "https://www.nepa.unicamp.br/arquivo/uploads/taco-4a-edicao/taco-4a-edicao-2/"
SHA1 = "4e6e03dd7baa94184496f4bf84905d17bb4444aa"  # planilha conferida em 30/09/2026
XLSX = ROOT / "data" / ".cache" / "taco-4a-edicao.xlsx"
OUT = ROOT / "data" / "taco.json"

if not XLSX.exists():
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    print("baixando a planilha da TACO...")
    req = urllib.request.Request(URL, headers={"User-Agent": "playbook-wr (scripts/taco.py)"})
    XLSX.write_bytes(urllib.request.urlopen(req, timeout=60).read())
sha = hashlib.sha1(XLSX.read_bytes()).hexdigest()
if sha != SHA1:
    sys.exit(f"A planilha mudou (sha1 {sha}). Confira as colunas antes de atualizar o SHA1 em scripts/taco.py.")

def num(v):
    """'Tr' (traço), 'NA' (não se aplica) e '*' (não analisado) viram 0."""
    return float(v) if isinstance(v, (int, float)) else 0.0

ws = load_workbook(XLSX, read_only=True, data_only=True)["CMVCol taco3"]
cats, foods, skipped, cat = [], [], [], None
for row in ws.iter_rows(values_only=True):
    n, desc = row[0], row[1]
    if isinstance(n, str) and desc is None and n.strip() and not n.strip().startswith(("Número", "Alimento")):
        cat = n.strip()
        continue
    if isinstance(n, (int, float)) and isinstance(desc, str):
        if row[3] == "*":  # energia não analisada na TACO: fica de fora em vez de aparecer com 0 kcal
            skipped.append(desc)
            continue
        if cat not in cats:
            cats.append(cat)
        # colunas: 3 energia (kcal), 5 proteína, 6 lipídeos, 8 carboidrato (g por 100 g)
        foods.append([" ".join(desc.split()), cats.index(cat), round(num(row[3])), round(num(row[5]), 1), round(num(row[8]), 1), round(num(row[6]), 1)])

names = [f[0] for f in foods]
assert len(foods) > 550 and len(set(names)) == len(names), (len(foods), len(set(names)))
json.dump({"fonte": "TACO 4ª edição (NEPA/UNICAMP, 2011)", "cats": cats, "foods": foods},
          open(OUT, "w", encoding="utf-8", newline="\n"), ensure_ascii=False, indent=0)
print(f"TACO: {len(foods)} alimentos em {len(cats)} grupos ({len(skipped)} sem energia analisada ficaram de fora) -> {OUT}")
