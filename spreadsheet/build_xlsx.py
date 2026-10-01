import math, os, sys, json
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "data"))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from foods import FOODS, TREINO, SABADO, DESCANSO
from units import UNITS

# Perfil: lido de spreadsheet/perfil.local.json (fica fora do Git). Sem o arquivo, usa valores neutros.
PERFIL = {"peso": 75, "altura": 175, "idade": 25, "sexo": "M", "obj": "Manter"}
_local = Path(__file__).resolve().parent / "perfil.local.json"
if _local.exists():
    PERFIL.update(json.loads(_local.read_text(encoding="utf-8")))
_fmt_n = lambda x: f"{x:g}".replace(".", ",")
PERFIL_TXT = f"{_fmt_n(PERFIL['peso'])} kg, {PERFIL['altura']} cm, {PERFIL['idade']} anos, objetivo {PERFIL['obj']}"

FN = "Arial"
NAVY, BLUE_BAND, GREY = "1F3864", "D9E2F3", "F2F2F2"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HDR_FILL = PatternFill("solid", fgColor=NAVY)
BAND_FILL = PatternFill("solid", fgColor=BLUE_BAND)
GREY_FILL = PatternFill("solid", fgColor=GREY)
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def f(bold=False, size=10, color="000000", italic=False):
    return Font(name=FN, bold=bold, size=size, color=color, italic=italic)

F_BASE = f(); F_BOLD = f(True); F_IN = f(color="0000FF"); F_LINK = f(color="008000")
F_HDR = f(True, 10, "FFFFFF"); F_TITLE = f(True, 16, NAVY); F_SUB = f(False, 10, "595959", True)

WRAP = Alignment(wrap_text=True, vertical="center")
WRAPT = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

APP = {"gym": {}, "campo": {}, "sat": [], "aq": []}
wb = Workbook()
wb.remove(wb.active)

def sheet(name, tab, widths):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws

def title(ws, text, sub, last_col):
    ws["A1"] = text; ws["A1"].font = F_TITLE
    ws.merge_cells(f"A1:{last_col}1"); ws.row_dimensions[1].height = 26
    if sub:
        ws["A2"] = sub; ws["A2"].font = F_SUB; ws["A2"].alignment = WRAP
        ws.merge_cells(f"A2:{last_col}2")

def hdr(ws, row, labels, start=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=t)
        c.font = F_HDR; c.fill = HDR_FILL; c.alignment = CENTER; c.border = BORDER
    ws.row_dimensions[row].height = 30

def band(ws, row, text, last_col):
    ws[f"A{row}"] = text; ws[f"A{row}"].font = F_BOLD
    for col in range(1, ord(last_col) - 64 + 1):
        ws.cell(row=row, column=col).fill = BAND_FILL
    ws.merge_cells(f"A{row}:{last_col}{row}")
    ws.row_dimensions[row].height = 20

def note(ws, row, text, last_col, italic=False, bold=False):
    ws[f"A{row}"] = text; ws[f"A{row}"].font = f(bold, 10, "000000", italic); ws[f"A{row}"].alignment = WRAPT
    ws.merge_cells(f"A{row}:{last_col}{row}")

def inp(c, fmt=None):
    c.font = F_IN; c.fill = INPUT_FILL; c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if fmt: c.number_format = fmt

def fit_rows(ws, line_h=13.0):
    merged = {}
    for mr in ws.merged_cells.ranges:
        w = sum((ws.column_dimensions[chr(64 + c)].width or 8.43) for c in range(mr.min_col, mr.max_col + 1))
        merged[(mr.min_row, mr.min_col)] = w
    for row in ws.iter_rows():
        mx = 1
        for c in row:
            v = c.value
            if isinstance(v, str) and not v.startswith("="):
                w = merged.get((c.row, c.column), ws.column_dimensions[chr(64 + c.column)].width or 8.43)
                cpl = max(int(w * 1.05), 4)
                lines = sum(max(1, math.ceil(len(p) / cpl)) for p in v.split("\n"))
                mx = max(mx, lines)
        cur = ws.row_dimensions[row[0].row].height
        need = mx * line_h + 4
        if mx > 1 or cur is None:
            ws.row_dimensions[row[0].row].height = max(need, cur or 0, 18)

# =====================================================================
# 1) INÍCIO
# =====================================================================
ws = sheet("Início", "1F3864", [26, 70])
title(ws, "Plano WR – Treino & Alimentação (Futebol Americano)",
      "Wide Receiver: velocidade, aceleração, mudança de direção, mãos firmes e prevenção de lesões. Treino do time: sábado, 12h–16h.", "B")
band(ws, 4, "Como usar em 3 passos", "B")
steps = [
 ("1) Preencha o Perfil", "Aba 'Perfil': coloque seu peso, altura, idade, sexo e objetivo. As metas de calorias e macros de cada tipo de dia se calculam sozinhas."),
 ("2) Siga a Semana", "Aba 'Semana': visão geral de Seg a Dom (academia, campo, time e descanso) + a meta de alimentação de cada dia."),
 ("3) Registre e evolua", "Anote as cargas nas abas Academia A/B/C, suas marcas em 'Progresso' e ajuste as porções nos Cardápios."),
]
r = 5
for a, b in steps:
    ws[f"A{r}"] = a; ws[f"A{r}"].font = F_BOLD; ws[f"A{r}"].alignment = WRAP
    ws[f"B{r}"] = b; ws[f"B{r}"].font = F_BASE; ws[f"B{r}"].alignment = WRAP
    r += 1
band(ws, 9, "Legenda de cores", "B")
c = ws["A10"]; c.value = "Texto azul / fundo amarelo"; inp(c)
ws["B10"] = "Células que VOCÊ edita (peso, cargas, quantidades, marcas)."; ws["B10"].font = F_BASE
ws["A11"] = "Texto preto"; ws["A11"].font = F_BASE; ws["A11"].alignment = CENTER
ws["B11"] = "Fórmulas e textos fixos – não precisa mexer."; ws["B11"].font = F_BASE
ws["A12"] = "Texto verde"; ws["A12"].font = F_LINK; ws["A12"].alignment = CENTER
ws["B12"] = "Valores puxados de outra aba (ex.: metas do Perfil)."; ws["B12"].font = F_BASE
band(ws, 14, "Abas da planilha (clique para ir)", "B")
tabs = [
 ("Perfil", "Seus dados + metas diárias de calorias, proteína, carboidrato e gordura."),
 ("Semana", "Rotina Seg–Dom: o que fazer em cada dia e a meta nutricional do dia."),
 ("Academia A", "Segunda – Inferior (força + posterior de coxa). Registro de cargas por semana."),
 ("Academia B", "Quarta – Superior + core + pegada."),
 ("Academia C", "Quinta – Potência, saltos e prevenção de lesões."),
 ("Campo", "Terça (velocidade/aceleração) e Sexta (mãos, rotas e agilidade)."),
 ("Sábado (Time)", "Rotina do dia do treino do time: aquecimento, horários, hidratação e recuperação."),
 ("Cardápio Treino", "Cardápio de Seg–Sex com quantidades e macros calculados."),
 ("Cardápio Sábado", "Cardápio pensado para o treino de 12h às 16h (pré, durante e pós)."),
 ("Cardápio Descanso", "Cardápio de domingo (recuperação)."),
 ("Base Alimentos", "Tabela de calorias/macros por 100 g – você pode adicionar alimentos."),
 ("Progresso", "Registro semanal: peso, sono, testes de velocidade/salto e cargas-chave."),
 ("Dicas", "Progressão, deload, hidratação, sono, suplementos e sinais de alerta."),
]
r = 15
for name, desc in tabs:
    c = ws[f"A{r}"]; c.value = name; c.hyperlink = f"#'{name}'!A1"
    c.font = Font(name=FN, size=10, color="0563C1", underline="single")
    ws[f"B{r}"] = desc; ws[f"B{r}"].font = F_BASE; ws[f"B{r}"].alignment = WRAP
    r += 1
r += 1
band(ws, r, "Avisos importantes", "B"); r += 1
avisos = [
 f"Perfil atual: {PERFIL_TXT}. Se algo estiver diferente, corrija na aba Perfil e as metas se atualizam.",
 "Plano baseado em práticas comuns de preparação física e nutrição esportiva; não substitui acompanhamento de nutricionista, preparador físico ou médico. Se tiver dor, lesão ou condição de saúde, procure um profissional antes de seguir.",
 "Valores nutricionais são aproximados (ordem de grandeza da Tabela TACO/rótulos) e variam por marca e preparo.",
]
for a in avisos:
    ws[f"A{r}"] = a; ws[f"A{r}"].font = F_BASE; ws[f"A{r}"].alignment = WRAPT
    ws.merge_cells(f"A{r}:B{r}"); r += 1
fit_rows(ws)

# =====================================================================
# 2) PERFIL
# =====================================================================
ws = sheet("Perfil", "C00000", [34, 16, 16, 14, 14, 14, 14, 14, 14, 40])
title(ws, "Perfil e Metas Diárias", f"Edite as células amarelas. Tudo abaixo delas se recalcula. Dados atuais: {PERFIL_TXT}.", "J")
band(ws, 4, "1) Seus dados", "J")
rows = [(5, "Peso (kg)", PERFIL["peso"], "0.0", "Atualize se o peso mudar"),
        (6, "Altura (cm)", PERFIL["altura"], "0", "Em centímetros"),
        (7, "Idade (anos)", PERFIL["idade"], "0", "Atualize no seu aniversário"),
        (8, "Sexo (M/F)", PERFIL["sexo"], None, "Usado na fórmula do metabolismo basal"),
        (9, "Objetivo", PERFIL["obj"], None, "Definir com manutenção de força (lista: Ganhar massa / Manter / Definir)"),
        (10, "Proteína (g por kg)", 2.0, "0.0", "Em déficit, 2,0–2,4 g/kg ajuda a preservar músculo e força")]
for r, lab, val, fmt, nt in rows:
    ws[f"A{r}"] = lab; ws[f"A{r}"].font = F_BOLD
    ws[f"B{r}"] = val; inp(ws[f"B{r}"], fmt)
    ws[f"C{r}"] = nt; ws[f"C{r}"].font = F_SUB
    ws.merge_cells(f"C{r}:J{r}")
dv1 = DataValidation(type="list", formula1='"M,F"', allow_blank=False); ws.add_data_validation(dv1); dv1.add("B8")
dv2 = DataValidation(type="list", formula1='"Ganhar massa,Manter,Definir"', allow_blank=False); ws.add_data_validation(dv2); dv2.add("B9")
band(ws, 12, "2) Ajuste calórico por objetivo (multiplicador sobre a manutenção)", "J")
for r, lab, val in [(13, "Ganhar massa", 1.10), (14, "Manter", 1.00), (15, "Definir", 0.90)]:
    ws[f"A{r}"] = lab; ws[f"A{r}"].font = F_BASE
    ws[f"B{r}"] = val; inp(ws[f"B{r}"], "0.00")
ws["C13"] = "+10% ≈ ganho lento (0,25–0,5 kg/semana). Definir: −10% ≈ perda de ~0,3–0,7 kg/semana, preservando a força."; ws["C13"].font = F_SUB
ws.merge_cells("C13:J13")
ws["A16"] = "Ajuste aplicado"; ws["A16"].font = F_BOLD
ws["B16"] = "=INDEX(B13:B15,MATCH(B9,A13:A15,0))"; ws["B16"].font = F_BASE; ws["B16"].number_format = "0.00"; ws["B16"].alignment = CENTER
band(ws, 18, "3) Metabolismo basal (Mifflin-St Jeor)", "J")
ws["A19"] = "TMB (kcal/dia em repouso)"; ws["A19"].font = F_BOLD
ws["B19"] = '=10*B5+6.25*B6-5*B7+IF(B8="M",5,-161)'; ws["B19"].font = F_BASE; ws["B19"].number_format = "#,##0"; ws["B19"].alignment = CENTER
band(ws, 21, "4) Metas diárias por tipo de dia", "J")
hdr(ws, 22, ["Tipo de dia", "Fator de atividade", "Manutenção (kcal)", "META (kcal)", "Proteína (g)", "Carbo (g/kg)", "Carbo (g)", "Gordura (g)", "Gordura (g/kg)", "Checagem"])
days = [(23, "Dia de treino (Seg–Sex)", 1.55, 3.3), (24, "Sábado – treino do time 12h–16h", 1.75, 4.3), (25, "Domingo – descanso", 1.40, 2.6)]
for r, lab, fa, cg in days:
    ws[f"A{r}"] = lab; ws[f"A{r}"].font = F_BOLD; ws[f"A{r}"].border = BORDER; ws[f"A{r}"].alignment = WRAP
    ws[f"B{r}"] = fa; inp(ws[f"B{r}"], "0.00")
    ws[f"C{r}"] = f"=$B$19*B{r}"
    ws[f"D{r}"] = f"=ROUND(C{r}*$B$16,-1)"
    ws[f"E{r}"] = "=ROUND($B$5*$B$10,0)"
    ws[f"F{r}"] = cg; inp(ws[f"F{r}"], "0.0")
    ws[f"G{r}"] = f"=ROUND($B$5*F{r},0)"
    ws[f"H{r}"] = f"=ROUND((D{r}-E{r}*4-G{r}*4)/9,0)"
    ws[f"I{r}"] = f"=H{r}/$B$5"
    ws[f"J{r}"] = f'=IF(I{r}<0.8,"Gordura baixa: reduza carbo ou suba kcal","OK")'
    for col, fmt in zip("CDEGHIJ", ["#,##0", "#,##0", "0", "0", "0", "0.00", "@"]):
        cc = ws[f"{col}{r}"]; cc.font = F_BOLD if col == "D" else F_BASE; cc.number_format = fmt
        cc.alignment = CENTER; cc.border = BORDER
    ws.row_dimensions[r].height = 30
band(ws, 27, "Como as metas são calculadas (premissas)", "J")
notes = [
 "• Manutenção = TMB × fator de atividade. Fatores usados (estimativas): 1,55 (treino diário moderado-alto), 1,75 (4 h de treino de time no sábado), 1,40 (descanso). Se seu peso não responder em 2 semanas, ajuste o fator.",
 "• Meta = manutenção × ajuste do objetivo. Definir (−10%): perda esperada de ~0,3–0,7 kg/semana. Se perder mais rápido que ~1% do peso por semana ou a força/rendimento cair, suba ~150–200 kcal (≈ +40 g de carboidrato). Se o peso estagnar por 2–3 semanas, reduza ~100–150 kcal ou some passos diários. (Para ganhar massa: se estagnar 2 semanas, some 150–200 kcal.)",
 "• Proteína fixa em g/kg do peso (faixa 1,6–2,2). Carboidrato varia por dia (mais no sábado, por causa das 4 h de treino). Gordura = restante das calorias (mínimo recomendado ~0,8 g/kg – a coluna 'Checagem' avisa).",
 "• Kcal por grama: proteína 4, carboidrato 4, gordura 9. Fórmula do metabolismo basal: Mifflin-St Jeor.",
 "• Como você está em déficit, os carboidratos foram mantidos o mais altos possível (principalmente no sábado) e a gordura fica perto do mínimo de ~0,8 g/kg. Estas são estimativas de partida: use a aba Progresso (peso semanal e força) para calibrar.",
]
for i, t in enumerate(notes):
    note(ws, 28 + i, t, "J")
fit_rows(ws)
ws.freeze_panes = "A4"

# =====================================================================
# 3) SEMANA
# =====================================================================
ws = sheet("Semana", "ED7D31", [12, 34, 16, 22, 14, 16, 20, 12, 11, 11, 11])
title(ws, "Rotina Semanal – Wide Receiver", "Sábado é o dia fixo do time (12h–16h). De segunda a sexta você treina à tarde (sugestão: 16h30): faça o pré-treino às 15h e a refeição pós-treino até 1–2 h depois.", "K")
hdr(ws, 4, ["Dia", "Foco do dia", "Onde", "Horário sugerido", "Duração", "Intensidade", "Cardápio", "Meta kcal", "Prot (g)", "Carbo (g)", "Gord (g)"])
week = [
 ("Segunda", "Academia A – Inferior (força + posterior de coxa)", "Academia", "Tarde (ex.: 16h30)", "60–75 min", "Alta", "Cardápio Treino", 23),
 ("Terça", "Campo – Velocidade e aceleração (sprints)", "Campo / quadra", "Tarde (ex.: 16h30)", "60 min", "Alta (qualidade)", "Cardápio Treino", 23),
 ("Quarta", "Academia B – Superior + core + pegada", "Academia", "Tarde (ex.: 16h30)", "60–70 min", "Moderada-alta", "Cardápio Treino", 23),
 ("Quarta (opcional)", "Flag football à noite – só se estiver bem", "Campo / quadra", "20:00", "cerca de 1 h", "Moderada", "", 0),
 ("Quinta", "Academia C – Potência, saltos e prevenção", "Academia", "Tarde (ex.: 16h30)", "60 min", "Moderada-alta", "Cardápio Treino", 23),
 ("Sexta", "Campo leve – Mãos, rotas e agilidade", "Campo / quadra", "Tarde (ex.: 16h30)", "45–60 min", "Baixa-moderada", "Cardápio Treino", 23),
 ("Sábado", "TREINO DO TIME (aquecimento individual antes)", "Campo do time", "12:00–16:00", "4 h", "Alta", "Cardápio Sábado", 24),
 ("Domingo", "Descanso ativo: caminhada leve, mobilidade 15 min, sono", "Livre", "Livre", "20–30 min", "Baixa", "Cardápio Descanso", 25),
]
r = 5
for d, foco, onde, hora, dur, inten, card, prow in week:
    vals = [d, foco, onde, hora, dur, inten, card]
    for i, v in enumerate(vals):
        c = ws.cell(row=r, column=1 + i, value=v); c.font = F_BOLD if i < 2 else F_BASE
        c.alignment = WRAP; c.border = BORDER
        if d == "Sábado": c.fill = PatternFill("solid", fgColor="FCE4D6")
        if prow == 0: c.fill = PatternFill("solid", fgColor="FFF2CC")
    if card != "":
        ws.cell(row=r, column=7).hyperlink = f"#'{card}'!A1"
        ws.cell(row=r, column=7).font = Font(name=FN, size=10, color="0563C1", underline="single")
    for j, col in enumerate("DEGH"):
        src = f"Perfil!{col}{prow}"
        c = ws.cell(row=r, column=8 + j, value=("—" if prow == 0 else f"={src}"))
        if prow == 0: c.fill = PatternFill("solid", fgColor="FFF2CC")
        c.font = F_LINK; c.number_format = "#,##0" if j == 0 else "0"; c.alignment = CENTER; c.border = BORDER
        if d == "Sábado": c.fill = PatternFill("solid", fgColor="FCE4D6")
    ws.row_dimensions[r].height = 32
    r += 1
band(ws, 14, "Regras de ouro da semana", "K")
rules = [
 "1. Sábado é o dia mais pesado: na sexta NÃO faça pernas pesadas (por isso a sexta é leve e técnica). Durma 8–9 h na sexta à noite.",
 "2. Segunda é o dia pós-time: se as pernas estiverem muito pesadas/doloridas, faça as mesmas séries com cargas ~10% menores e 1 série a menos (tudo bem, o objetivo é evoluir sem se machucar).",
 "3. Sprints e potência (terça e quinta) sempre com descanso completo entre séries: o treino é de qualidade, não de cansaço. Se a velocidade cair visivelmente, encerre o bloco.",
 "4. Semana 4 de cada ciclo = descarga (deload): reduza cargas ~10% e 1 série por exercício. Depois recomece o ciclo com cargas um pouco maiores.",
 "5. Alimentação: bata a meta de proteína todos os dias; carboidrato é o combustível dos dias de treino (principalmente sábado). Beba água ao longo do dia (urina clara/amarelo-claro).",
 "6. Você está em déficit (definir): mantenha as cargas pesadas nos exercícios principais (é o que preserva a força). Se a recuperação piorar, corte séries antes de cortar carga. Não adicione cardio extra intenso: os sprints e o time já cobrem; caminhadas leves (8–10 mil passos/dia) ajudam.",
 "7. Se estiver doente, com dor aguda ou dormindo pouco por vários dias, troque o treino por mobilidade/caminhada. Consistência ao longo dos meses vale mais que uma semana perfeita.",
 "8. Quarta às 20h tem flag football (opcional). O jantar de 19:15 fica perto do jogo: faça algo leve com carboidrato às 18:30 (ex.: 1 pão francês e 1 banana) e deixe o jantar completo para depois. Se jogar, na quinta corte 1 série dos saltos e do terra caso as pernas estejam pesadas.",
]
for i, t in enumerate(rules):
    note(ws, 15 + i, t, "K")
fit_rows(ws)
ws.freeze_panes = "A5"

# =====================================================================
# 4) ACADEMIA A/B/C
# =====================================================================
GYM_W = [4, 34, 8, 12, 11, 8, 10, 10, 10, 12, 62]
def gym(name, tab, ttl, sub, warm, exs, final):
    ws = sheet(name, tab, GYM_W)
    APP["gym"][name] = dict(title=ttl, sub=sub, warm=warm, exs=exs, final=final)
    title(ws, ttl, sub, "K")
    note(ws, 3, "Objetivo em déficit: MANTER (ou subir devagar) as cargas. Quando fizer TODAS as séries no topo da faixa de repetições com técnica limpa, some 2,5 kg (membros superiores: 1–2 kg) na semana seguinte; se a carga cair 2 semanas seguidas, revise sono/calorias. RPE = esforço percebido (10 = limite; 8 = sobrariam ~2 repetições). Semana 4 = descarga: cargas ~10% menores e 1 série a menos.", "K", italic=True)
    note(ws, 4, "Como preencher: anote na coluna da semana a carga (kg) usada nas séries de trabalho. Exemplo: Agachamento, Sem 1 = 60. Use 'PC' para peso do corpo.", "K", italic=True)
    hdr(ws, 6, ["#", "Exercício", "Séries", "Repetições", "Descanso", "RPE alvo", "Sem 1 (kg)", "Sem 2 (kg)", "Sem 3 (kg)", "Sem 4 deload (kg)", "Por que importa pro WR / dica de execução"])
    ws["A7"] = "AQ"; ws["A7"].font = F_BOLD; ws["A7"].alignment = CENTER
    ws["B7"] = warm; ws["B7"].font = f(False, 10, "000000", True); ws["B7"].alignment = WRAP
    ws.merge_cells("B7:K7")
    for col in range(1, 12): ws.cell(row=7, column=col).fill = GREY_FILL
    r = 8
    for i, (ex, s, rep, desc, rpe, tip) in enumerate(exs, 1):
        vals = [i, ex, s, rep, desc, rpe]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=1 + j, value=v); c.font = F_BOLD if j == 1 else F_BASE
            c.alignment = WRAP if j == 1 else CENTER; c.border = BORDER
        for col in range(7, 11):
            inp(ws.cell(row=r, column=col), "0.0")
        c = ws.cell(row=r, column=11, value=tip); c.font = F_BASE; c.alignment = WRAP; c.border = BORDER
        r += 1
    ws.cell(row=r, column=1, value="FIM"); ws.cell(row=r, column=1).font = F_BOLD; ws.cell(row=r, column=1).alignment = CENTER
    ws.cell(row=r, column=2, value=final).font = f(False, 10, "000000", True)
    ws.cell(row=r, column=2).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=11)
    for col in range(1, 12): ws.cell(row=r, column=col).fill = GREY_FILL
    fit_rows(ws)
    ws.freeze_panes = "C7"

gym("Academia A", "2E75B6", "Academia A – Inferior (segunda-feira)",
    "Foco: força de perna, cadeia posterior (isquiotibiais/glúteos) e estabilidade. Duração: 60–75 min. Volta do fim de semana: comece leve se estiver dolorido.",
    "Aquecimento (10 min): 5 min de bike/trote leve + mobilidade de quadril e tornozelo + 2 séries de agachamento com peso do corpo + 2×10 saltinhos rápidos (pogos).",
    [("Agachamento livre (back squat)", 4, "5", "2–3 min", "7–8", "Base da aceleração. Desça controlado, suba explosivo; calcanhares no chão e tronco firme. Se não tiver barra/rack, use agachamento goblet com halter pesado (4×8)."),
     ("Levantamento terra romeno (RDL)", 3, "8", "2 min", "7–8", "Fortalece isquiotibiais e glúteos – os músculos que mais se lesionam em sprints. Quadril para trás, costas neutras, barra colada nas pernas."),
     ("Afundo búlgaro (pé de trás no banco)", 3, "8 por perna", "90 s", "7–8", "Correr é uma ação de uma perna só: corrige assimetrias e dá estabilidade nos cortes."),
     ("Hip thrust", 3, "10", "90 s", "8", "Extensão de quadril = potência de saída. Pausa de 1 s no topo, sem arquear a lombar."),
     ("Nórdico (Nordic hamstring curl)", 3, "5–6", "2 min", "8–9", "Um dos exercícios com melhor evidência para prevenir lesão de isquiotibial. Desça o mais devagar possível; use as mãos para ajudar a subir."),
     ("Panturrilha em pé (elevação)", 4, "12–15", "60 s", "8", "Tendão de Aquiles e tornozelo fortes para frear e acelerar. Pausa de 1 s embaixo e no topo."),
     ("Pallof press (anti-rotação)", 3, "10 por lado", "60 s", "7", "Core que resiste à rotação: você recebe contato e muda de direção sem perder a postura.")],
    "Finalização (5 min): alongamento leve de quadril, quadríceps, isquiotibiais e panturrilha.")

gym("Academia B", "2E75B6", "Academia B – Superior + core + pegada (quarta-feira)",
    "Foco: força de empurrar/puxar, ombros saudáveis, pegada e core. Duração: 60–70 min.",
    "Aquecimento (8–10 min): 5 min de remo/corda + rotações de ombro + band pull-apart 2×15 + 1×10 flexões.",
    [("Supino reto (barra ou halteres)", 4, "6", "2–3 min", "7–8", "Força de empurrar: bloqueios e enfrentar o contato do defensor."),
     ("Remada curvada (barra ou halteres)", 4, "8", "2 min", "8", "Costas fortes protegem o ombro e melhoram o uso das mãos na saída da linha (release)."),
     ("Desenvolvimento com halteres", 3, "8", "90 s", "7–8", "Ombro forte e estável para bolas altas e contato."),
     ("Barra fixa (pull-up)", 3, "6–8 (ou máx.)", "2 min", "8–9", "Força relativa ao peso corporal, essencial para quem precisa ser rápido. Use elástico de assistência se necessário."),
     ("Face pull", 3, "15", "60 s", "7", "Saúde do ombro e postura (manguito e deltoide posterior)."),
     ("Rosca martelo + tríceps na corda (bi-set)", 3, "12 + 12", "60 s", "7–8", "Braços fortes ajudam a proteger a bola e nos contatos. Volume complementar."),
     ("Rosca de punho e rosca de punho reversa", 3, "15", "45 s", "8", "Antebraço e pegada: mãos firmes que seguram a bola sob contato."),
     ("Farmer's walk (caminhada do fazendeiro)", 3, "30 m", "60 s", "8", "Pegada + core + estabilidade do tronco."),
     ("Roda abdominal (ab wheel) ou dead bug", 3, "8–10", "60 s", "7", "Core anti-extensão: mantém a pelve estável durante sprints.")],
    "Finalização (5 min): alongar peitoral, dorsais e ombros.")

gym("Academia C", "2E75B6", "Academia C – Potência, saltos e prevenção (quinta-feira)",
    "Foco: explosão, mudança de direção e prevenção de lesões (isquiotibiais, virilha, tornozelo). Duração: ~60 min. Sexta é leve: não leve nada até a falha hoje.",
    "Aquecimento (10–12 min): 5 min de bike + mobilidade de quadril/tornozelo + 2×5 pogos + 2×3 saltos verticais submáximos.",
    [("Salto vertical (countermovement jump) – esforço máximo", 4, "4", "90 s", "Máx.", "Potência para pegar bolas no ponto mais alto e sair do chão. Qualidade > quantidade: pouso silencioso, joelhos alinhados."),
     ("Salto em distância parado (standing broad jump)", 3, "3", "90 s", "Máx.", "Potência horizontal = aceleração nos primeiros passos. Meça e anote a distância (app: Evolução > Testes de campo; planilha: aba Progresso)."),
     ("Kettlebell swing pesado (ou hang power clean, se tiver técnica/supervisão)", 4, "8 (swing) / 3 (clean)", "2 min", "7–8", "Potência de quadril. O hang clean só com técnica ensinada por treinador."),
     ("Levantamento terra convencional (barra reta)", 3, "5", "2–3 min", "7", "Força total da cadeia posterior. Carga moderada (deixe ~2–3 reps de reserva) para não acumular fadiga antes do sábado; costas neutras e barra colada nas pernas."),
     ("Step-up com halteres (caixote alto)", 3, "8 por perna", "90 s", "7–8", "Força unilateral e controle de joelho."),
     ("Salto lateral (skater jump)", 3, "6 por lado", "60 s", "Máx.", "Mudança de direção: absorver o impacto e empurrar lateralmente."),
     ("Arremesso de medicine ball (rotacional ou peito)", 3, "6", "60 s", "Máx.", "Potência de tronco e braços; ótimo para explosão."),
     ("Copenhagen plank (adutores)", 3, "20 s por lado", "60 s", "7", "Prevenção de lesão de virilha, comum em esportes com cortes."),
     ("Elevação de tibial anterior (encostado na parede)", 2, "15", "45 s", "6", "Canelas e tornozelo: ajudam na prevenção de dor na canela e no apoio da corrida.")],
    "Finalização (5–8 min): mobilidade de quadril/tornozelo. Sexta é leve – chegue descansado ao sábado.")

# =====================================================================
# 5) CAMPO
# =====================================================================
ws = sheet("Campo", "70AD47", [22, 40, 20, 14, 14, 62, 9])
title(ws, "Treinos de Campo – Velocidade, Rotas e Mãos", "Terça: velocidade/aceleração (alta intensidade, qualidade máxima). Sexta: mãos, rotas e agilidade (baixa-moderada, para chegar descansado ao sábado).", "G")
def campo_block(start, ttl, rows):
    APP["campo"][ttl] = rows
    band(ws, start, ttl, "G")
    hdr(ws, start + 1, ["Bloco", "Exercício / drill", "Séries × distância/reps", "Descanso", "Intensidade", "Como fazer / foco", "Feito?"])
    r = start + 2
    for row in rows:
        for j, v in enumerate(row):
            c = ws.cell(row=r, column=1 + j, value=v); c.font = F_BOLD if j == 1 else F_BASE
            c.alignment = CENTER if j in (2, 3, 4) else WRAP; c.border = BORDER
        inp(ws.cell(row=r, column=7)); ws.cell(row=r, column=7).border = BORDER
        r += 1
    return r
r = campo_block(4, "TERÇA – Velocidade e aceleração (60 min)", [
 ("Aquecimento", "Trote leve + mobilidade dinâmica (leg swings, círculos de quadril, afundo andando)", "8–10 min", "—", "Leve", "Suba a temperatura aos poucos. Tornozelos, quadril e isquiotibiais bem soltos antes de correr forte."),
 ("Técnica de corrida", "A-skip, B-skip, high knees, butt kicks", "2 × 20 jardas cada", "Voltar andando", "Leve", "Postura alta, apoio na bola do pé, cotovelos a 90°. Movimentos rápidos e precisos."),
 ("Ativação", "Strides progressivos", "3 × 30 jardas", "Voltar andando", "60% → 90%", "Acelere gradualmente, sem forçar. Prepara o corpo para os sprints máximos."),
 ("Aceleração", "Saídas de 10 jardas (posição de 3 apoios ou 2 apoios)", "6 × 10 jardas", "60–90 s", "95–100%", "Tronco inclinado ~45°, passos curtos e potentes, empurre o chão para trás. Alterne o pé da frente."),
 ("Velocidade", "Sprint de 20 jardas", "4 × 20 jardas", "2 min", "95–100%", "Transição da aceleração para a postura ereta; mantenha relaxamento no rosto e ombros."),
 ("Velocidade máx.", "Sprint de 40 jardas (o teste clássico do combine)", "3 × 40 jardas", "3 min", "95–100%", "Frequência alta de passada, ombros relaxados. Se perder qualidade, encerre o bloco."),
 ("Resistido", "Sprint com elástico, trenó ou subida leve", "4 × 15 jardas", "2 min", "90%", "Carga leve (trenó ≤ ~10% do peso corporal) para não distorcer a técnica. Foco na projeção do corpo."),
 ("Volta à calma", "Trote leve + alongamento de quadril, isquiotibiais e panturrilha", "10 min", "—", "Leve", "Respiração calma; desacelere gradualmente."),
])
r = campo_block(r + 1, "SEXTA – Mãos, rotas e agilidade (45–60 min, intensidade controlada)", [
 ("Aquecimento", "Trote + mobilidade dinâmica + 2 strides", "10 min", "—", "Leve", "Sem cansar: hoje é dia de técnica, não de volume."),
 ("Pés rápidos", "Escada de agilidade (4 padrões: 1 apoio, 2 apoios, Ickey shuffle, lateral)", "2 × cada padrão", "30 s", "Moderada", "Pés leves, olhar para frente. Qualidade de cada toque."),
 ("Agilidade", "5-10-5 (pro agility)", "4 (2 p/ cada lado)", "90 s", "85%", "Baixe o centro de gravidade, toque a linha com a mão e explode para o lado oposto."),
 ("Agilidade", "3-cone drill (L-drill)", "4", "90 s", "85%", "Curvas fechadas, cabeça e ombros liderando a curva. Anote o tempo (app: Evolução > Testes de campo; planilha: aba Progresso)."),
 ("Saída de break", "Plant & drive (frear e sair) em cone", "3 por lado", "60 s", "80–90%", "Corte em 3 passos, sem arredondar. Firme o pé de fora e empurre o chão."),
 ("Rotas", "Árvore de rotas: slant, out, curl, comeback, in/dig, post, corner, go", "2 × cada rota", "Voltar andando", "70–85%", "Com QB/parceiro: stem vertical agressivo, venda a rota (ombros/cabeça), quebra limpa, cabeça e olhos para a bola ao sair do break."),
 ("Mãos", "Bola de tênis contra parede (mão alternada)", "3 × 30 s", "30 s", "Moderada", "Reação e coordenação olho-mão; mãos suaves."),
 ("Recepção", "Catch drills: bola alta (high point), baixa, por cima do ombro, lateral (toe-tap na linha)", "4 × 5 de cada tipo", "45 s", "Moderada", "Olhos na bola até guardá-la (tuck), mãos em diamante, pegue com as mãos e não com o corpo. Proteja a bola após a recepção."),
 ("Pós-recepção", "Recepção + corte + sprint curto (YAC – jardas após a recepção)", "4 × 15 jardas", "60 s", "85%", "Guarde a bola no braço de fora, mude de direção e acelere."),
 ("Mobilidade", "Alongamento de quadril, isquiotibiais, panturrilha e ombros", "10 min", "—", "Leve", "Prepare o corpo para o sábado: relaxe e hidrate-se."),
])
fit_rows(ws)
ws.freeze_panes = "A4"

# =====================================================================
# 6) SÁBADO (TIME)
# =====================================================================
ws = sheet("Sábado (Time)", "FFC000", [16, 46, 62])
title(ws, "Sábado – Dia do Treino do Time (12h às 16h)", "Cronograma sugerido de alimentação, hidratação, aquecimento e recuperação em torno das 4 horas de treino.", "C")
hdr(ws, 4, ["Horário", "O que fazer", "Nutrição / hidratação"])
sat = [
 ("Sexta à noite", "Durma 8–9 h. Evite treino pesado de perna na sexta.", "Jantar com carboidrato (arroz, macarrão, batata). Beba água ao longo do dia; nada de álcool."),
 ("07:30", "Acordar com calma; café da manhã.", "Cardápio Sábado: café da manhã (~1/5 das calorias do dia). Beba ~500 ml de água."),
 ("09:30", "Refeição pré-treino principal (2h30 antes).", "Carboidrato + proteína magra, pouca gordura e pouca fibra (evita desconforto). Não teste alimentos novos no dia do treino."),
 ("11:00", "Chegar ao campo. Aquecimento individual (12–15 min), ver quadro abaixo.", "Beba 300–500 ml de água/isotônico até o início."),
 ("11:15", "Lanche leve (banana + mel + isotônico).", "Carboidrato rápido 30–45 min antes."),
 ("12:00–16:00", "TREINO DO TIME. Nas pausas, hidrate-se e volte a ativar (pequenas corridas) para não esfriar.", "Água/isotônico em toda pausa (200–250 ml a cada 15–20 min). Carboidrato leve ao longo do treino: banana, isotônico, água de coco."),
 ("16:00–16:30", "Volta à calma: trote leve 5 min + alongamento/mobilidade 5–10 min.", "Continue bebendo. Pese-se antes e depois: para cada 1 kg perdido, reponha ~1,2–1,5 L de líquido nas horas seguintes."),
 ("16:30", "Refeição pós-treino (até 1 h depois).", "Carboidrato + proteína (ex.: arroz, feijão, carne) para recuperar glicogênio e músculo. Ver Cardápio Sábado."),
 ("19:45", "Jantar leve a moderado.", "Proteína + carboidrato + vegetais."),
 ("22:00", "Ceia e dormir cedo.", "Iogurte + aveia + castanhas. Meta: 8–9 h de sono."),
]
APP["sat"] = sat
r = 5
for t, a, b in sat:
    for j, v in enumerate([t, a, b]):
        c = ws.cell(row=r, column=1 + j, value=v); c.font = F_BOLD if j == 0 else F_BASE
        c.alignment = WRAP; c.border = BORDER
        if t == "12:00–16:00":
            c.fill = PatternFill("solid", fgColor="FCE4D6")
    r += 1
r += 1
band(ws, r, "Aquecimento individual do WR (12–15 min antes do time)", "C"); r += 1
hdr(ws, r, ["Ordem", "Exercício", "Detalhes"]); r += 1
aq = [
 (1, "Trote leve", "3 min, ritmo confortável."),
 (2, "Mobilidade dinâmica", "Leg swings frontal e lateral (10 cada perna), círculos de quadril, afundo com rotação (6 por lado)."),
 (3, "Técnica de corrida", "A-skip e B-skip 2 × 20 m; carioca/shuffle lateral 2 × 20 m."),
 (4, "Ativação de glúteo e isquiotibial", "Ponte de glúteo com 1 perna (10 por lado) ou caminhada com mini band (10 passos por lado)."),
 (5, "Strides progressivos", "3 × 30 m, de 60% a 90%."),
 (6, "Específico de WR", "2 acelerações de 10 jardas + 2 saídas de break (plant & drive) + 2 recepções fáceis."),
]
APP["aq"] = aq
for n, ex, det in aq:
    for j, v in enumerate([n, ex, det]):
        c = ws.cell(row=r, column=1 + j, value=v); c.font = F_BOLD if j == 1 else F_BASE
        c.alignment = CENTER if j == 0 else WRAP; c.border = BORDER
    r += 1
fit_rows(ws)
ws.freeze_panes = "A5"

# =====================================================================
# 7) BASE ALIMENTOS (antes dos cardápios porque eles referenciam)
# =====================================================================
wsb = sheet("Base Alimentos", "7F7F7F", [40, 12, 12, 12, 12, 42, 22, 24, 14])
title(wsb, "Base de Alimentos (valores por 100 g ou 100 ml)", "Valores aproximados (ordem de grandeza da Tabela TACO e rótulos comuns). Confira o rótulo do seu produto. As colunas G a I definem a medida caseira usada nos cardápios (sem balança). Para incluir um alimento novo, escreva nas linhas em branco abaixo (até a linha 40).", "I")
hdr(wsb, 4, ["Alimento", "kcal", "Proteína (g)", "Carbo (g)", "Gordura (g)", "Medida caseira", "Unidade (singular)", "Unidade (plural)", "g ou ml por unidade"])
for i, (n, k, p, c_, g, m) in enumerate(FOODS):
    r = 5 + i
    us, up, ug, st, self_, short = UNITS[n]
    ulabel_s = us if self_ else f"{us} de {short}"
    ulabel_p = up if self_ else f"{up} de {short}"
    for j, v in enumerate([n, k, p, c_, g, m, ulabel_s, ulabel_p, ug]):
        cell = wsb.cell(row=r, column=1 + j, value=v); cell.font = F_BASE if j == 0 else F_IN
        cell.border = BORDER; cell.alignment = WRAP if j in (0, 5, 6, 7) else CENTER
        if j in (1, 2, 3, 4): cell.number_format = "0.0" if j > 1 else "0"
for r in range(5 + len(FOODS), 41):
    for j in range(1, 10):
        cell = wsb.cell(row=r, column=j); cell.fill = INPUT_FILL; cell.font = F_IN; cell.border = BORDER
wsb.freeze_panes = "A5"

# =====================================================================
# 8) CARDÁPIOS
# =====================================================================
def cardapio(name, tab, ttl, sub, meals, prow, tips):
    ws = sheet(name, tab, [30, 13, 40, 11, 10, 10, 10, 10, 34, 34])
    title(ws, ttl, sub, "J")
    hdr(ws, 4, ["Refeição", "Horário", "Alimento (escolha na lista)", "Qtd (g ou ml)", "kcal", "Prot (g)", "Carbo (g)", "Gord (g)", "Medida caseira (referência)", "Em medidas caseiras (sem balança)"])
    dv = DataValidation(type="list", formula1="='Base Alimentos'!$A$5:$A$40", allow_blank=True)
    dv.error = "Escolha um alimento da lista ou adicione-o antes na aba 'Base Alimentos'."
    dv.errorTitle = "Alimento não encontrado"
    ws.add_data_validation(dv)
    r = 5; headers = []
    for mname, mtime, foods in meals:
        h = r; headers.append(h)
        ws[f"A{h}"] = mname; ws[f"B{h}"] = mtime
        for col in range(1, 11):
            ws.cell(row=h, column=col).fill = BAND_FILL; ws.cell(row=h, column=col).border = BORDER
            ws.cell(row=h, column=col).font = F_BOLD
        ws[f"A{h}"].alignment = WRAP; ws[f"B{h}"].alignment = CENTER
        last = h + len(foods) + 1  # inclui 1 linha vazia
        for col in "EFGH":
            ws[f"{col}{h}"] = f"=SUM({col}{h+1}:{col}{last})"
            ws[f"{col}{h}"].number_format = "#,##0" if col == "E" else "0"; ws[f"{col}{h}"].alignment = CENTER
        rr = h + 1
        for (fn, q) in foods + [(None, None)]:
            ws[f"C{rr}"] = fn; ws[f"D{rr}"] = q
            inp(ws[f"C{rr}"]); ws[f"C{rr}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            inp(ws[f"D{rr}"], "0")
            dv.add(f"C{rr}")
            for col, idx in zip("EFGH", "BCDE"):
                ws[f"{col}{rr}"] = (f'=IFERROR(IF($C{rr}="","",$D{rr}/100*INDEX(\'Base Alimentos\'!${idx}$5:${idx}$40,'
                                    f'MATCH($C{rr},\'Base Alimentos\'!$A$5:$A$40,0))),"")')
                ws[f"{col}{rr}"].font = F_LINK; ws[f"{col}{rr}"].border = BORDER; ws[f"{col}{rr}"].alignment = CENTER
                ws[f"{col}{rr}"].number_format = "#,##0" if col == "E" else "0.0"
            ws[f"I{rr}"] = (f'=IFERROR(IF($C{rr}="","",INDEX(\'Base Alimentos\'!$F$5:$F$40,MATCH($C{rr},\'Base Alimentos\'!$A$5:$A$40,0))),"")')
            ws[f"I{rr}"].font = F_LINK; ws[f"I{rr}"].border = BORDER; ws[f"I{rr}"].alignment = WRAP
            q = f"ROUND($D{rr}/INDEX('Base Alimentos'!$I$5:$I$40,MATCH($C{rr},'Base Alimentos'!$A$5:$A$40,0))*2,0)/2"
            ws[f"J{rr}"] = (f'=IFERROR(IF(OR($C{rr}="",$D{rr}=""),"",{q}&" "&IF({q}>1,'
                            f"INDEX('Base Alimentos'!$H$5:$H$40,MATCH($C{rr},'Base Alimentos'!$A$5:$A$40,0)),"
                            f"INDEX('Base Alimentos'!$G$5:$G$40,MATCH($C{rr},'Base Alimentos'!$A$5:$A$40,0)))),\"\")")
            ws[f"J{rr}"].font = F_LINK; ws[f"J{rr}"].border = BORDER; ws[f"J{rr}"].alignment = WRAP
            for col in "AB": ws[f"{col}{rr}"].border = BORDER
            rr += 1
        r = rr
    r += 1
    T, M, Dd, P, Fc = r, r + 1, r + 2, r + 3, r + 4
    labels = {T: "TOTAL PLANEJADO", M: "META DO DIA (aba Perfil)", Dd: "DIFERENÇA (plano − meta)", P: "% DA META", Fc: "FATOR DE AJUSTE DAS PORÇÕES"}
    for rowx, lab in labels.items():
        ws[f"A{rowx}"] = lab; ws[f"A{rowx}"].font = F_BOLD
        ws.merge_cells(f"A{rowx}:D{rowx}")
        for col in range(1, 11):
            ws.cell(row=rowx, column=col).border = BORDER
            if rowx == T: ws.cell(row=rowx, column=col).fill = BAND_FILL
    for col, pc in zip("EFGH", "DEGH"):
        ws[f"{col}{T}"] = "=" + "+".join(f"{col}{h}" for h in headers)
        ws[f"{col}{M}"] = f"=Perfil!${pc}${prow}"
        ws[f"{col}{Dd}"] = f"={col}{T}-{col}{M}"
        ws[f"{col}{P}"] = f'=IF({col}{M}=0,"",{col}{T}/{col}{M})'
        fmt = "#,##0" if col == "E" else "0"
        ws[f"{col}{T}"].number_format = fmt; ws[f"{col}{T}"].font = F_BOLD
        ws[f"{col}{M}"].number_format = fmt; ws[f"{col}{M}"].font = F_LINK
        ws[f"{col}{Dd}"].number_format = "+#,##0;-#,##0;0" if col == "E" else "+0;-0;0"; ws[f"{col}{Dd}"].font = F_BASE
        ws[f"{col}{P}"].number_format = "0%"; ws[f"{col}{P}"].font = F_BASE
        for rowx in (T, M, Dd, P): ws[f"{col}{rowx}"].alignment = CENTER
    ws[f"E{Fc}"] = f'=IF(E{T}=0,"",E{M}/E{T})'; ws[f"E{Fc}"].number_format = "0.00"; ws[f"E{Fc}"].font = F_BOLD; ws[f"E{Fc}"].alignment = CENTER
    ws[f"F{Fc}"] = "← multiplique as quantidades por este fator para bater a meta de calorias (ex.: 1,05 = 5% a mais)."
    ws[f"F{Fc}"].font = F_SUB; ws[f"F{Fc}"].alignment = WRAP
    ws.merge_cells(f"F{Fc}:J{Fc}")
    green = PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE"); red = PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")
    rng = f"E{P}:H{P}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["0.95", "1.05"], fill=green))
    ws.conditional_formatting.add(rng, CellIsRule(operator="notBetween", formula=["0.9", "1.1"], fill=red))
    rr = Fc + 2
    band(ws, rr, "Dicas deste cardápio", "J"); rr += 1
    for t in tips:
        note(ws, rr, t, "J"); rr += 1
    note(ws, rr, "As quantidades são um ponto de partida. Se as calorias do cardápio ficarem longe das metas do Perfil, use o fator de ajuste ou troque alimentos (escolha na lista da coluna C; para incluir novos, use a aba Base Alimentos). Alimentos equivalentes: frango ↔ tilápia/atum/carne magra; arroz ↔ batata/macarrão/mandioca. A coluna J converte gramas em medidas caseiras (arredondadas para meia unidade), então dá para seguir sem balança.", "J", italic=True)
    fit_rows(ws)
    ws.freeze_panes = "A5"

cardapio("Cardápio Treino", "5B9BD5", "Cardápio – Dias de Treino (Segunda a Sexta)",
         "Pensado para treinar às ~16h30: almoço às 12h30, lanche pré-treino às 15h e refeição completa depois do treino. Se mudar o horário, mova só o pré-treino e o pós.",
         TREINO, 23,
         ["• Distribua a proteína em 4–6 refeições (aprox. 25–40 g cada) para otimizar a recuperação muscular.",
          "• Pré-treino: carboidrato de fácil digestão e pouca gordura/fibra. Pós-treino: carboidrato + proteína.",
          "• Na sexta, mantenha o cardápio: os carboidratos ajudam a chegar com o estoque de energia cheio para o sábado. Em déficit, não corte os carboidratos ao redor do treino: corte primeiro doces e frituras."])
cardapio("Cardápio Sábado", "FFC000", "Cardápio – Sábado (treino do time 12h–16h)",
         "Estratégia: café da manhã + refeição principal 2h30 antes + lanche rápido + carboidrato durante + refeição completa depois.",
         SABADO, 24,
         ["• A refeição das 9h30 é a mais importante para o desempenho: carboidrato + proteína magra, pouca gordura e fibra.",
          "• Durante o treino (12h–16h): 200–250 ml de líquido a cada 15–20 min e cerca de 30–60 g de carboidrato por hora (isotônico + banana). Ajuste conforme o suor e o calor.",
          "• Pós-treino: refeição completa até 1 h depois. Se não tiver fome, comece por líquidos + carboidrato e complete em seguida.",
          "• Não experimente alimentos novos no dia do treino."])
cardapio("Cardápio Descanso", "A5A5A5", "Cardápio – Domingo (descanso e recuperação)",
         "Menos carboidrato que o sábado, mantendo a proteína alta para reparar o músculo.",
         DESCANSO, 25,
         ["• Priorize alimentos frescos, vegetais e boa hidratação. É um bom dia para preparar as marmitas da semana.",
          "• Termine o dia cedo: o sono é o principal recurso de recuperação."])

# =====================================================================
# 9) PROGRESSO
# =====================================================================
ws = sheet("Progresso", "7030A0", [10, 12, 11, 11, 10, 10, 12, 12, 13, 12, 14, 14, 36])
title(ws, "Progresso Semanal", "Pese-se sempre no mesmo dia/horário (ex.: domingo de manhã, em jejum). Faça os testes de velocidade/salto a cada 4 semanas. Linha cinza = EXEMPLO (apague ou ignore).", "M")
hdr(ws, 4, ["Semana", "Data", "Peso (kg)", "Δ peso (kg)", "Sono médio (h)", "Energia (1–5)", "10 jardas (s)", "40 jardas (s)", "Salto vertical (cm)", "5-10-5 (s)", "Agachamento 5RM (kg)", "Supino 6RM (kg)", "Observações"])
ex = ["Ex.", "06/10", float(PERFIL["peso"]), None, 7.5, 4, 1.75, 4.85, 55, 4.40, 90, 70, "Exemplo – apague ou ignore"]
for j, v in enumerate(ex):
    c = ws.cell(row=5, column=1 + j, value=v); c.font = f(False, 10, "7F7F7F", True); c.fill = GREY_FILL; c.border = BORDER; c.alignment = CENTER
for w in range(1, 13):
    r = 5 + w
    ws.cell(row=r, column=1, value=w).font = F_BOLD
    ws.cell(row=r, column=1).alignment = CENTER; ws.cell(row=r, column=1).border = BORDER
    for col in range(2, 14):
        if col == 4:
            c = ws.cell(row=r, column=4)
            if w > 1:
                c.value = f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(C{r-1})),C{r}-C{r-1},"")'
            c.font = F_BASE; c.number_format = "+0.0;-0.0;0.0"; c.alignment = CENTER; c.border = BORDER
        else:
            c = ws.cell(row=r, column=col); inp(c, "0.00" if col in (7, 8, 10) else ("0.0" if col in (3, 5) else None))
            if col == 13: c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
note(ws, 19, "Como interpretar (objetivo Definir): perda ideal de ~0,3–0,7 kg/semana com a força mantida (agachamento e supino estáveis). Se perder mais rápido que ~1% do peso/semana, ou os sprints/cargas caírem, suba ~150–200 kcal/dia. Se o peso estagnar por 2–3 semanas, reduza ~100–150 kcal ou aumente os passos diários. Dormir mais costuma ser o primeiro ajuste de desempenho.", "M", italic=True)
fit_rows(ws)
ws.freeze_panes = "B5"

# =====================================================================
# 10) DICAS
# =====================================================================
ws = sheet("Dicas", "00B0F0", [28, 110])
title(ws, "Dicas, Progressão e Cuidados", "Guia rápido para tirar o máximo do plano com segurança.", "B")
hdr(ws, 4, ["Tema", "Orientação"])
tips = [
 ("Progressão de carga", "Some 2,5 kg (parte superior: 1–2 kg) quando completar todas as séries no topo da faixa de repetições com boa técnica. Se travar em duas semanas, mantenha a carga e busque mais 1–2 repetições."),
 ("Descarga (deload)", "A cada 4 semanas, reduza as cargas em ~10% e faça 1 série a menos por exercício. Recuperação também faz o atleta evoluir."),
 ("Sono", "Meta: 8–9 h por noite, principalmente na sexta (véspera do time) e no domingo. Horário regular e celular fora da cama ajudam."),
 ("Hidratação", "Beba água ao longo do dia; urina amarelo-claro é um bom sinal. No sábado, siga o plano da aba 'Sábado (Time)' e pese-se antes e depois."),
 ("Creatina (opcional)", "A creatina monoidratada (3–5 g por dia, todos os dias) é o suplemento com maior evidência para força e potência em adultos saudáveis. Converse com um profissional de saúde antes de usar."),
 ("Whey protein (opcional)", "É apenas comodidade para bater a meta de proteína; a comida de verdade vem primeiro. Confira o rótulo (30 g de produto ≠ 30 g de proteína)."),
 ("Cafeína (opcional)", "Pode melhorar o desempenho, mas teste primeiro em treinos comuns (nunca no dia do time pela primeira vez) e evite à noite, para não atrapalhar o sono."),
 ("Prevenção de lesões", "Pontos críticos para WR: isquiotibiais (nórdico, RDL), virilha/adutores (Copenhagen), tornozelo (panturrilha, tibial) e ombro (face pull). Não pule o aquecimento."),
 ("Mobilidade diária (10 min)", "Quadril (90/90, afundo com rotação), tornozelo, isquiotibiais e ombro. Pode fazer à noite, antes de dormir."),
 ("Ajuste de calorias (definir)", "Perda de ~0,3–0,7 kg/semana. Se perder mais rápido que ~1% do peso por semana ou a força cair, some 150–200 kcal (≈ +40 g de carboidrato). Se estagnar por 2–3 semanas, reduza ~100–150 kcal ou aumente os passos diários. Mantenha o déficit por blocos de 8–12 semanas e faça uma semana de manutenção depois."),
 ("Definir sem perder força", "Proteína alta (2,0–2,4 g/kg), cargas pesadas mantidas nos básicos (agachamento, supino, remada, terra), déficit moderado, sono de 8 h e carboidratos concentrados ao redor do treino. Evite déficits agressivos: em esporte de velocidade, o rendimento cai rápido."),
 ("Sinais para parar e procurar ajuda", "Dor aguda ou articular, formigamento, tontura, dor no peito, falta de ar fora do normal ou queda persistente do rendimento e do humor. Interrompa o treino e procure um profissional de saúde."),
 ("Nível intermediário", "Com 6 meses a 2 anos de treino, a progressão linear (somar carga toda semana) funciona por algumas semanas; em déficit, o realista é manter as cargas. Se estagnar, varie a faixa de repetições (ex.: 4×5 → 3×6) antes de mexer no volume total."),
 ("Consistência", "Um plano executado 80% do tempo por meses vence um plano perfeito por duas semanas. Se perder um treino, siga para o próximo — não tente 'compensar' dobrando."),
]
for i, (a, b) in enumerate(tips):
    r = 5 + i
    ws.cell(row=r, column=1, value=a).font = F_BOLD
    ws.cell(row=r, column=2, value=b).font = F_BASE
    for col in (1, 2):
        ws.cell(row=r, column=col).alignment = WRAP; ws.cell(row=r, column=col).border = BORDER
note(ws, 5 + len(tips) + 1, "Este material é educativo e não substitui acompanhamento de nutricionista, preparador físico ou médico.", "B", italic=True)
fit_rows(ws)
ws.freeze_panes = "A5"

# Ordem final das abas
order = ["Início", "Perfil", "Semana", "Academia A", "Academia B", "Academia C", "Campo", "Sábado (Time)",
         "Cardápio Treino", "Cardápio Sábado", "Cardápio Descanso", "Base Alimentos", "Progresso", "Dicas"]
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.properties.title = "Plano WR – Treino e Alimentação"
OUT = ROOT / "dist"
OUT.mkdir(exist_ok=True)
wb.save(OUT / "Plano_WR_Treino_e_Alimentacao.xlsx")
print("planilha:", OUT / "Plano_WR_Treino_e_Alimentacao.xlsx")

# conteúdo dos treinos (Academia A/B/C, campo, sábado, aquecimento) usado pelo app
json.dump(APP, open(ROOT / "data" / "training.json", "w", encoding="utf-8", newline="\n"), ensure_ascii=False)
print("treinos:", ROOT / "data" / "training.json")
